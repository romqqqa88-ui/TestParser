from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from pandas import DatetimeTZDtype
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from avito_parser_console.export.naming import build_export_filename


class ExcelExporter:
    def export(
        self,
        rows: list[dict],
        export_dir: str,
        columns: list[str] | None = None,
        header_map: dict[str, str] | None = None,
        query_tag: str | None = None,
    ) -> str:
        path = Path(export_dir)
        path.mkdir(parents=True, exist_ok=True)
        filename = build_export_filename("avito_export", "xlsx", query_tag=query_tag)
        full_path = path / filename

        df = pd.DataFrame(rows)
        for column in df.columns:
            if isinstance(df[column].dtype, DatetimeTZDtype):
                df[column] = df[column].dt.tz_localize(None)
            elif df[column].dtype == "object":
                df[column] = df[column].map(
                    lambda value: value.replace(tzinfo=None) if isinstance(value, datetime) and value.tzinfo else value
                )
        if columns:
            df = df[columns]
        if header_map:
            df = df.rename(columns=header_map)
        df.to_excel(full_path, index=False)
        from openpyxl import load_workbook

        book = load_workbook(full_path)
        ws = book.active
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for idx, col in enumerate(ws.columns, 1):
            max_length = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(max_length + 2, 12), 60)
        book.save(full_path)
        return str(full_path)
